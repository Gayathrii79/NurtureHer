"""Dataset loading utilities for the PCOS Random Forest training pipeline.

Two supported input formats:

1. ``simple`` -- the flat NurtureHer schema::

       age,bmi,cycle_irregularity,hair_growth,skin_darkening,weight_gain,follicle_count,pcos

2. ``kaggle`` -- the public Kaggle dataset *Polycystic ovary syndrome (PCOS)*
   by Prasoon Kottarathil (541 patients from 10 hospitals in Kerala, India),
   distributed as ``PCOS_data_without_infertility.xlsx`` (sheet ``Full_new``)
   or community CSV conversions of the same file. Column headers are matched
   after normalization, so variants such as ``Follicle No. (L)`` and
   ``Follicle.No..L.`` both resolve.

Both loaders emit features in :data:`app.ml.features.FEATURE_ORDER` so the
trained model stays aligned with the live prediction endpoint.
"""

from __future__ import annotations

import csv
import logging
import re
from pathlib import Path

from app.ml.features import FEATURE_ORDER

logger = logging.getLogger(__name__)

_MISSING_TOKENS = {"", "#NAME?", "#VALUE!", "#DIV/0!", "NA", "N/A", "NAN", "NONE", "NULL", "-"}

# Normalized header -> canonical field name for the Kaggle dataset.
KAGGLE_HEADER_MAP = {
    "pcosyn": "label",
    "ageyrs": "age",
    "age": "age",
    "weightkg": "weight",
    "heightcm": "height",
    "heightcms": "height",
    "bmi": "bmi",
    "cycleri": "cycle",
    "weightgainyn": "weight_gain",
    "hairgrowthyn": "hair_growth",
    "skindarkeningyn": "skin_darkening",
    "folliclenol": "follicle_left",
    "folliclenor": "follicle_right",
}

KAGGLE_REQUIRED = {
    "label",
    "age",
    "cycle",
    "weight_gain",
    "hair_growth",
    "skin_darkening",
    "follicle_left",
    "follicle_right",
}


def parse_bool(value: object) -> float:
    return 1.0 if str(value).strip().lower() in {"1", "1.0", "true", "yes", "y"} else 0.0


def normalize_header(name: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(name).lower())


def _to_float(value: object) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if text.upper() in _MISSING_TOKENS:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _read_rows(path: Path) -> list[dict[str, object]]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - depends on optional dep
            raise RuntimeError(
                "Install openpyxl to read the Kaggle .xlsx dataset (pip install openpyxl), "
                "or fetch the CSV mirror via `python -m scripts.fetch_pcos_dataset`."
            ) from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook["Full_new"] if "Full_new" in workbook.sheetnames else workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(cell) if cell is not None else "" for cell in next(iterator)]
        rows = [dict(zip(headers, row)) for row in iterator]
        workbook.close()
        return rows
    with path.open(newline="", encoding="utf-8-sig") as csv_file:
        return list(csv.DictReader(csv_file))


def detect_format(headers) -> str:
    normalized = {normalize_header(header) for header in headers}
    if "pcosyn" in normalized:
        return "kaggle"
    if "pcos" in normalized:
        return "simple"
    raise ValueError(
        "Unrecognized dataset headers; expected the NurtureHer simple schema or the Kaggle PCOS dataset"
    )


def _load_simple_rows(rows: list[dict[str, object]], label_column: str = "pcos") -> tuple[list[list[float]], list[int]]:
    features: list[list[float]] = []
    labels: list[int] = []
    for row in rows:
        features.append(
            [
                float(row["age"]),
                float(row["bmi"]),
                parse_bool(row["cycle_irregularity"]),
                parse_bool(row["hair_growth"]),
                parse_bool(row["skin_darkening"]),
                parse_bool(row["weight_gain"]),
                float(row.get("follicle_count") or 0),
            ]
        )
        labels.append(int(row[label_column]))
    return features, labels


def _load_kaggle_rows(rows: list[dict[str, object]]) -> tuple[list[list[float]], list[int]]:
    if not rows:
        raise ValueError("Dataset is empty")

    header_lookup: dict[str, str] = {}
    for raw_header in rows[0].keys():
        canonical = KAGGLE_HEADER_MAP.get(normalize_header(raw_header))
        if canonical and canonical not in header_lookup:
            header_lookup[canonical] = str(raw_header)

    missing = KAGGLE_REQUIRED - set(header_lookup)
    if missing:
        raise ValueError(f"Kaggle dataset is missing expected columns: {sorted(missing)}")

    def field(row: dict[str, object], name: str) -> object:
        header = header_lookup.get(name)
        return row.get(header) if header else None

    features: list[list[float]] = []
    labels: list[int] = []
    skipped = 0
    for row in rows:
        label = _to_float(field(row, "label"))
        age = _to_float(field(row, "age"))
        bmi = _to_float(field(row, "bmi"))
        if bmi is None:
            # Some CSV conversions carry spreadsheet artifacts (e.g. "#NAME?") in
            # the BMI column; recompute from weight/height when possible.
            weight = _to_float(field(row, "weight"))
            height = _to_float(field(row, "height"))
            if weight and height:
                bmi = round(weight / ((height / 100) ** 2), 2)
        cycle = _to_float(field(row, "cycle"))
        follicles = [
            value
            for value in (_to_float(field(row, "follicle_left")), _to_float(field(row, "follicle_right")))
            if value is not None
        ]
        if label is None or age is None or bmi is None or cycle is None or not follicles:
            skipped += 1
            continue
        features.append(
            [
                age,
                bmi,
                1.0 if cycle >= 3 else 0.0,  # Kaggle encoding: 2 = regular, 4 = irregular
                1.0 if (_to_float(field(row, "hair_growth")) or 0) >= 1 else 0.0,
                1.0 if (_to_float(field(row, "skin_darkening")) or 0) >= 1 else 0.0,
                1.0 if (_to_float(field(row, "weight_gain")) or 0) >= 1 else 0.0,
                max(follicles),
            ]
        )
        labels.append(int(label))

    if skipped:
        logger.info("Skipped %s Kaggle rows with unusable values", skipped)
    if not features:
        raise ValueError("No usable rows found in the Kaggle dataset")
    return features, labels


def load_training_csv(path: str | Path, label_column: str = "pcos") -> tuple[list[list[float]], list[int]]:
    """Backwards-compatible loader for the simple NurtureHer CSV schema."""
    return _load_simple_rows(_read_rows(Path(path)), label_column)


def load_dataset(path: str | Path, dataset_format: str = "auto") -> tuple[list[list[float]], list[int], str]:
    """Load a PCOS training dataset, returning ``(features, labels, detected_format)``."""
    dataset_path = Path(path)
    rows = _read_rows(dataset_path)
    if not rows:
        raise ValueError(f"No rows found in {dataset_path}")
    fmt = detect_format(rows[0].keys()) if dataset_format == "auto" else dataset_format
    if fmt == "kaggle":
        features, labels = _load_kaggle_rows(rows)
    elif fmt == "simple":
        features, labels = _load_simple_rows(rows)
    else:
        raise ValueError(f"Unknown dataset format: {dataset_format}")
    logger.info("Loaded %s samples (%s positive) from %s [%s format]", len(labels), sum(labels), dataset_path, fmt)
    return features, labels, fmt


def feature_names() -> list[str]:
    return list(FEATURE_ORDER)
